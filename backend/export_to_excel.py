"""Export all database tables into a single Excel workbook (.xlsx) with styled sheets.

Usage:
    # Export local SQLite database:
    python export_to_excel.py

    # Or export live Cloud SQL PostgreSQL database:
    python export_to_excel.py --url "postgresql+psycopg://user:pass@host:5432/kebabzilla"
"""

import argparse
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine, inspect, text


def main():
    parser = argparse.ArgumentParser(description="Export database tables to an Excel workbook")
    parser.add_argument(
        "--url",
        default=os.environ.get("DATABASE_URL", "sqlite:///./kebabzilla.db"),
        help="Database URL to export (default: sqlite:///./kebabzilla.db)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output Excel filename (default: KebabZilla_Database.xlsx on Desktop or current folder)",
    )
    args = parser.parse_args()

    db_url = args.url
    if db_url.startswith("postgres://"):
        db_url = db_url.replace("postgres://", "postgresql+psycopg://", 1)
    elif db_url.startswith("postgresql://") and not db_url.startswith("postgresql+"):
        db_url = db_url.replace("postgresql://", "postgresql+psycopg://", 1)

    print(f"[*] Connecting to database: {db_url.split('@')[-1] if '@' in db_url else db_url}")
    engine = create_engine(db_url)
    inspector = inspect(engine)
    table_names = sorted(inspector.get_table_names())

    if not table_names:
        print("[!] No tables found in database.")
        sys.exit(0)

    # Output path: Desktop or current directory
    output_path = args.output
    if not output_path:
        desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        if os.path.exists(desktop_dir):
            output_path = os.path.join(desktop_dir, "KebabZilla_Database.xlsx")
        else:
            output_path = "KebabZilla_Database.xlsx"

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid") # Dark slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    cell_alignment = Alignment(vertical="center")

    total_rows = 0
    with engine.connect() as conn:
        for table in table_names:
            # Excel sheet names can be max 31 characters
            sheet_title = table[:31]
            ws = wb.create_sheet(title=sheet_title)

            # Query table data
            result = conn.execute(text(f'SELECT * FROM "{table}"'))
            columns = list(result.keys())
            rows = result.fetchall()
            total_rows += len(rows)

            # Write header row
            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = thin_border
            ws.row_dimensions[1].height = 24

            # Write data rows
            for row_idx, row in enumerate(rows, 2):
                for col_idx, value in enumerate(row, 1):
                    # Convert dicts / lists to formatted JSON string
                    if isinstance(value, (dict, list)):
                        cell_value = json.dumps(value, ensure_ascii=False)
                    elif value is None:
                        cell_value = ""
                    else:
                        cell_value = str(value)

                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                    cell.font = cell_font
                    cell.alignment = cell_alignment
                    cell.border = thin_border
                ws.row_dimensions[row_idx].height = 20

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                # Clamp column width between 12 and 45 characters
                ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 45))

            # Enable Excel auto-filter on header row
            if columns:
                last_col = get_column_letter(len(columns))
                ws.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 1)}"

            print(f"[+] Exported table '{table}': {len(rows)} rows, {len(columns)} columns")

    wb.save(output_path)
    print(f"\n[SUCCESS] Excel workbook generated successfully!")
    print(f"File saved to: {output_path}")
    print(f"Total tables: {len(table_names)}, Total rows: {total_rows}")


if __name__ == "__main__":
    main()
